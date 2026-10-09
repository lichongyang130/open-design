from __future__ import annotations

import contextlib
import csv
import io
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

import sys
SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import reconcile_spreadsheets as reconcile_tool  # noqa: E402


class ReconciliationCsvTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-spreadsheet-reconcile-")
        self.root = Path(self.temp.name)
        self.file_a = self.root / "a.csv"
        self.file_b = self.root / "b.csv"
        self.output = self.root / "reconciliation.xlsx"
        self.report = self.root / "reconciliation.json"
        self.file_a.write_text(
            "OrderID,Amount,Status\n"
            "001,100,paid\n"
            "002,200,pending\n"
            "003,50,paid\n"
            "004,10,review\n"
            "004,11,review\n"
            ",7,blank-key\n",
            encoding="utf-8",
        )
        self.file_b.write_text(
            "OrderID,Amount,Status\n"
            "001,100.00,paid\n"
            "002,250,pending\n"
            "005,99,new\n"
            "004,10,review\n",
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def run_reconciliation(self) -> dict:
        return reconcile_tool.reconcile(
            self.file_a,
            self.file_b,
            self.output,
            self.report,
            key_columns=["OrderID"],
            compare_columns=["Amount", "Status"],
        )

    def test_matches_by_key_and_separates_differences_unmatched_duplicate_and_blank_keys(self) -> None:
        report = self.run_reconciliation()
        summary = report["summary"]

        self.assertEqual(summary["rows_in_a"], 6)
        self.assertEqual(summary["rows_in_b"], 4)
        self.assertEqual(summary["matched_unique_keys"], 2)
        self.assertEqual(summary["matched_without_differences"], 1)
        self.assertEqual(summary["matched_with_differences"], 1)
        self.assertEqual(summary["field_differences"], 1)
        self.assertEqual(summary["only_in_a"], 1)
        self.assertEqual(summary["only_in_b"], 1)
        self.assertEqual(summary["duplicate_keys"], 1)
        self.assertEqual(summary["rows_with_blank_keys_a"], 1)
        self.assertEqual(report["differences"][0]["column"], "Amount")
        self.assertEqual(report["differences"][0]["value_a"], "200")
        self.assertEqual(report["differences"][0]["value_b"], "250")
        self.assertEqual(report["only_in_a"][0]["key"]["OrderID"], "003")
        self.assertEqual(report["only_in_b"][0]["key"]["OrderID"], "005")
        self.assertEqual(report["duplicate_keys"][0]["action"], "not_auto_paired")
        self.assertEqual(report["needs_review"][0]["issue"], "blank_key")

    def test_numeric_text_compares_equal_but_leading_zero_keys_are_preserved(self) -> None:
        self.assertEqual(reconcile_tool.canonical("100"), reconcile_tool.canonical("100.00"))
        self.assertEqual(reconcile_tool.canonical("100"), reconcile_tool.canonical(100))
        self.assertNotEqual(reconcile_tool.canonical("00123"), reconcile_tool.canonical(123))

        report = self.run_reconciliation()
        self.assertEqual(report["summary"]["matched_without_differences"], 1)

    def test_cli_writes_validated_workbook_and_json_audit(self) -> None:
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            result = reconcile_tool.run([
                str(self.file_a),
                str(self.file_b),
                "--key", "OrderID",
                "--compare", "Amount",
                "--compare", "Status",
                "--output", str(self.output),
                "--report", str(self.report),
            ])

        self.assertEqual(result, 0)
        self.assertTrue(self.output.exists())
        self.assertTrue(self.report.exists())
        self.assertIn('"ok": true', stdout.getvalue().lower())
        import json
        report_data = json.loads(self.report.read_text(encoding="utf-8"))
        self.assertEqual(report_data["summary"]["field_differences"], 1)
        workbook = openpyxl.load_workbook(self.output, read_only=True, data_only=True)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Summary", "Field Differences", "Only in A", "Only in B", "Duplicate Keys", "Needs Review"],
            )
            self.assertEqual(workbook["Field Differences"].max_row, 2)
            self.assertEqual(workbook["Only in A"].max_row, 2)
            self.assertEqual(workbook["Only in B"].max_row, 2)
        finally:
            workbook.close()

    def test_rejects_input_file_as_output_path(self) -> None:
        with self.assertRaises(reconcile_tool.ReconciliationError):
            reconcile_tool.run([
                str(self.file_a),
                str(self.file_b),
                "--key", "OrderID",
                "--output", str(self.file_a),
                "--report", str(self.report),
            ])
        self.assertIn("OrderID", self.file_a.read_text(encoding="utf-8"))


@unittest.skipIf(openpyxl is None, "openpyxl is required for Excel reconciliation tests")
class ReconciliationExcelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-spreadsheet-reconcile-xlsx-")
        self.root = Path(self.temp.name)
        self.file_a = self.root / "a.xlsx"
        self.file_b = self.root / "b.xlsx"
        self.output = self.root / "reconciliation.xlsx"
        self.report = self.root / "reconciliation.json"

        for path, amount in ((self.file_a, 100), (self.file_b, 150)):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Orders"
            sheet.append(["OrderID", "Amount"])
            sheet.append(["A-001", amount])
            workbook.save(path)
            workbook.close()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_reads_named_worksheets_and_compares_excel_files(self) -> None:
        report = reconcile_tool.reconcile(
            self.file_a,
            self.file_b,
            self.output,
            self.report,
            key_columns=["OrderID"],
            compare_columns=["Amount"],
            sheet_a="Orders",
            sheet_b="Orders",
        )
        self.assertEqual(report["selected_sheet_a"], "Orders")
        self.assertEqual(report["selected_sheet_b"], "Orders")
        self.assertEqual(report["summary"]["matched_with_differences"], 1)
        self.assertEqual(report["differences"][0]["value_a"], 100)
        self.assertEqual(report["differences"][0]["value_b"], 150)


if __name__ == "__main__":
    unittest.main()
