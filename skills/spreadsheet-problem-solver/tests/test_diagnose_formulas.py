from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

SKILL_ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import diagnose_formulas as diagnostic  # noqa: E402


@unittest.skipIf(openpyxl is None, "openpyxl is required for formula diagnostic tests")
class FormulaDiagnosticTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-formula-diagnostics-")
        self.root = Path(self.temp.name)
        self.input_path = self.root / "sales.xlsx"
        self.output_path = self.root / "sales_formulas_repaired.xlsx"
        self.report_path = self.root / "formula_report.json"

    def tearDown(self) -> None:
        self.temp.cleanup()

    def make_gap_workbook(self, lower_formula: str = "=B4*2") -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Item", "Amount", "Double"])
        sheet.append(["A", 10, "=B2*2"])
        sheet.append(["B", 20, None])
        sheet.append(["C", 30, lower_formula])
        workbook.save(self.input_path)
        workbook.close()

    def diagnose(self, apply_safe_repairs: bool = False, sheets: list[str] | None = None) -> dict:
        report = diagnostic.build_report(
            self.input_path, self.output_path, dry_run=not apply_safe_repairs
        )
        return diagnostic.diagnose_workbook(
            self.input_path,
            self.output_path,
            report=report,
            selected_sheets=sheets,
            apply_safe_repairs=apply_safe_repairs,
        )

    def test_read_only_diagnosis_suggests_only_formula_confirmed_by_both_neighbors(self) -> None:
        self.make_gap_workbook()

        report = self.diagnose()

        candidates = [
            finding for finding in report["findings"]
            if finding["kind"] == "missing_formula_gap"
        ]
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["sheet"], "Sales")
        self.assertEqual(candidates[0]["cell"], "C3")
        self.assertEqual(candidates[0]["candidate_formula"], "=B3*2")
        self.assertFalse(self.output_path.exists())
        original = openpyxl.load_workbook(self.input_path, data_only=False)
        try:
            self.assertIsNone(original["Sales"]["C3"].value)
        finally:
            original.close()

    def test_apply_safe_repairs_writes_copy_and_preserves_source(self) -> None:
        self.make_gap_workbook()

        report = self.diagnose(apply_safe_repairs=True)

        self.assertTrue(self.output_path.exists())
        self.assertEqual(report["summary"]["safe_formula_gaps_repaired"], 1)
        output = openpyxl.load_workbook(self.output_path, data_only=False)
        source = openpyxl.load_workbook(self.input_path, data_only=False)
        try:
            self.assertEqual(output.sheetnames, ["Sales"])
            self.assertEqual(output["Sales"]["C3"].value, "=B3*2")
            self.assertIsNone(source["Sales"]["C3"].value)
        finally:
            output.close()
            source.close()

    def test_formula_mismatch_is_not_repaired(self) -> None:
        self.make_gap_workbook(lower_formula="=B4+2")

        report = self.diagnose()

        self.assertEqual(report["summary"]["safe_formula_gap_candidates"], 0)
        self.assertFalse(any(
            finding["kind"] == "missing_formula_gap"
            for finding in report["findings"]
        ))

    def test_reports_broken_references_and_literal_excel_error_values(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A1"] = "=SUM(#REF!)"
        sheet["A2"] = "#DIV/0!"
        workbook.save(self.input_path)
        workbook.close()

        report = self.diagnose()

        self.assertTrue(any(
            finding["kind"] == "broken_reference" and finding["cell"] == "A1"
            for finding in report["findings"]
        ))
        self.assertTrue(any(
            finding["kind"] == "excel_error_value" and finding["cell"] == "A2"
            for finding in report["findings"]
        ))

    def test_rejects_unknown_sheet_without_writing_output(self) -> None:
        self.make_gap_workbook()

        report = diagnostic.build_report(
            self.input_path, self.output_path, dry_run=True
        )
        with self.assertRaises(diagnostic.FormulaDiagnosticError):
            diagnostic.diagnose_workbook(
                self.input_path,
                self.output_path,
                report=report,
                selected_sheets=["Does Not Exist"],
            )
        self.assertFalse(self.output_path.exists())


if __name__ == "__main__":
    unittest.main()
