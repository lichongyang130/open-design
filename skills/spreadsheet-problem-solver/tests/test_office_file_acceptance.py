from __future__ import annotations

import contextlib
import csv
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import clean_spreadsheet as cleaner  # noqa: E402
import detect_anomalies as anomaly_tool  # noqa: E402
import diagnose_formulas as formula_tool  # noqa: E402
import reconcile_spreadsheets as reconcile_tool  # noqa: E402


class SalesOfficeFileAcceptanceTests(unittest.TestCase):
    """End-to-end acceptance on realistic temporary office files, not mocks."""

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-office-file-acceptance-")
        self.root = Path(self.temp.name)
        self.sales = self.root / "销售订单.csv"
        self.receipts = self.root / "回款记录.csv"
        self.cleaned = self.root / "销售订单_cleaned.csv"
        self.clean_report = self.root / "销售订单_清理报告.json"
        self.anomaly_report = self.root / "销售订单_异常报告.json"
        self.reconciliation_book = self.root / "订单回款_差异报告.xlsx"
        self.reconciliation_report = self.root / "订单回款_差异报告.json"
        self.formula_input = self.root / "订单计算.xlsx"
        self.formula_output = self.root / "订单计算_formulas_repaired.xlsx"
        self.formula_report = self.root / "订单计算_公式诊断.json"

        self.sales.write_text(
            "OrderID,Customer,Amount,Status,DoubleAmount\n"
            " A-001 ,Acme Ltd,100.00,paid,200.00\n"
            "A-002, Beta LLC ,250,pending,500\n"
            "A-003,Gamma Inc,50,paid,100\n"
            "A-004,Delta Co,9999,review,19998\n"
            "A-004,Delta Co,9999,review,19998\n"
            "A-006,,12,review,24\n"
            "\n",
            encoding="utf-8",
        )
        self.receipts.write_text(
            "OrderID,Amount,Status\n"
            "A-001,100,paid\n"
            "A-002,275,pending\n"
            "A-003,50,paid\n"
            "A-004,9999,review\n"
            "A-005,49,new\n",
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def invoke(self, fn, argv: list[str]) -> dict:
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            result = fn(argv)
        self.assertEqual(result, 0, stdout.getvalue())
        payload = json.loads(stdout.getvalue())
        self.assertTrue(payload["ok"])
        return payload

    def test_sales_workflows_produce_reviewable_files_and_preserve_sources(self) -> None:
        original_sales = self.sales.read_bytes()
        original_receipts = self.receipts.read_bytes()

        # 1. Run the real cleaner CLI: opt-in row removal and exact-duplicate
        # removal, with a standalone output file and an auditable JSON report.
        clean_payload = self.invoke(
            cleaner.run,
            [
                str(self.sales),
                "--drop-empty-rows",
                "--dedupe",
                "--output",
                str(self.cleaned),
                "--report",
                str(self.clean_report),
            ],
        )
        self.assertTrue(self.cleaned.is_file())
        clean_report = json.loads(self.clean_report.read_text(encoding="utf-8"))
        self.assertEqual(clean_report["summary"]["blank_rows_removed"], 1)
        self.assertEqual(clean_report["summary"]["duplicate_rows_removed"], 1)
        with self.cleaned.open("r", encoding="utf-8", newline="") as handle:
            cleaned_rows = list(csv.DictReader(handle))
        self.assertEqual(len(cleaned_rows), 5)
        self.assertEqual(cleaned_rows[1]["Customer"], "Beta LLC")

        # 2. Run the read-only anomaly audit on the original source and ensure
        # blank required values, duplicate records, and outlier candidates are
        # findings, not mutations.
        self.invoke(
            anomaly_tool.run,
            [
                str(self.sales),
                "--required-column", "Customer",
                "--numeric-column", "Amount",
                "--report", str(self.anomaly_report),
            ],
        )
        anomaly_report = json.loads(self.anomaly_report.read_text(encoding="utf-8"))
        self.assertGreaterEqual(anomaly_report["summary"]["missing_values"], 1)
        self.assertGreaterEqual(anomaly_report["summary"]["exact_duplicate_rows"], 1)
        self.assertGreaterEqual(anomaly_report["summary"]["numeric_outlier_candidates"], 1)

        # 3. Reconcile sales against payments by the explicit OrderID business
        # key. Duplicate A-004 stays isolated, amount conflicts are explicit,
        # and records missing from either file are named.
        self.invoke(
            reconcile_tool.run,
            [
                str(self.sales),
                str(self.receipts),
                "--key", "OrderID",
                "--compare", "Amount",
                "--compare", "Status",
                "--output", str(self.reconciliation_book),
                "--report", str(self.reconciliation_report),
            ],
        )
        reconciliation = json.loads(self.reconciliation_report.read_text(encoding="utf-8"))
        self.assertEqual(reconciliation["summary"]["field_differences"], 1)
        self.assertEqual(reconciliation["summary"]["only_in_a"], 1)
        self.assertEqual(reconciliation["summary"]["only_in_b"], 1)
        self.assertEqual(reconciliation["summary"]["duplicate_keys"], 1)
        self.assertTrue(self.reconciliation_book.is_file())
        report_book = openpyxl.load_workbook(self.reconciliation_book, read_only=True, data_only=True)
        try:
            self.assertEqual(
                report_book.sheetnames,
                ["Summary", "Field Differences", "Only in A", "Only in B", "Duplicate Keys", "Needs Review"],
            )
        finally:
            report_book.close()

        # 4. Diagnose and repair one missing formula only where the formulas
        # directly above and below independently imply the same formula.
        formula_book = openpyxl.Workbook()
        sheet = formula_book.active
        sheet.title = "订单计算"
        sheet.append(["OrderID", "Amount", "DoubleAmount"])
        sheet.append(["A-001", 100, "=B2*2"])
        sheet.append(["A-002", 250, None])
        sheet.append(["A-003", 50, "=B4*2"])
        sheet.append(["A-004", 9999, "=SUM(#REF!)"])
        formula_book.save(self.formula_input)
        formula_book.close()
        self.invoke(
            formula_tool.run,
            [
                str(self.formula_input),
                "--apply-safe-repairs",
                "--output", str(self.formula_output),
                "--report", str(self.formula_report),
            ],
        )
        formula_report = json.loads(self.formula_report.read_text(encoding="utf-8"))
        self.assertEqual(formula_report["summary"]["safe_formula_gaps_repaired"], 1)
        self.assertGreaterEqual(formula_report["summary"]["broken_reference_findings"], 1)
        fixed_book = openpyxl.load_workbook(self.formula_output, data_only=False)
        try:
            self.assertEqual(fixed_book["订单计算"]["C3"].value, "=B3*2")
            self.assertEqual(fixed_book["订单计算"]["C5"].value, "=SUM(#REF!)")
        finally:
            fixed_book.close()
        # All source files remain byte-for-byte intact.
        self.assertEqual(self.sales.read_bytes(), original_sales)
        self.assertEqual(self.receipts.read_bytes(), original_receipts)
        self.assertTrue(self.formula_input.is_file())
        self.assertTrue(clean_payload["ok"])


if __name__ == "__main__":
    unittest.main()
