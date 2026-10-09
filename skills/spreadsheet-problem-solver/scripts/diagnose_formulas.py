#!/usr/bin/env python3
"""Diagnose spreadsheet formula errors and optionally repair proven single-cell gaps.

This tool never overwrites the input workbook. Formula gaps are only considered
safe when translating the formula from the row above and the row below into the
blank cell produces exactly the same formula.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


MAX_CELLS_PER_SHEET = 1_000_000
MAX_RECORDED_ITEMS = 5_000
KNOWN_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#UNKNOWN!",
    "#GETTING_DATA",
}


class FormulaDiagnosticError(Exception):
    """Actionable input, safety, or validation failure."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def is_blank(value: Any) -> bool:
    return value is None or value == ""


def is_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() in KNOWN_ERROR_VALUES


def record(report: dict[str, Any], finding: dict[str, Any]) -> None:
    findings = report["findings"]
    if len(findings) < MAX_RECORDED_ITEMS:
        findings.append(finding)
    else:
        report["details_truncated"] = True


def formula_count(workbook: Any, worksheets: Iterable[Any] | None = None) -> int:
    count = 0
    selected = worksheets if worksheets is not None else workbook.worksheets
    for worksheet in selected:
        for row in worksheet.iter_rows():
            for cell in row:
                if is_formula(cell.value):
                    count += 1
    return count


def safe_gap_formula(worksheet: Any, row: int, column: int, translator: Any) -> str | None:
    """Return a candidate only when both adjacent formulas imply the same text."""
    from openpyxl.cell.cell import MergedCell

    target = worksheet.cell(row=row, column=column)
    above = worksheet.cell(row=row - 1, column=column)
    below = worksheet.cell(row=row + 1, column=column)
    if isinstance(target, MergedCell):
        return None
    if not is_blank(target.value) or target.comment is not None or target.hyperlink is not None:
        return None
    if not is_formula(above.value) or not is_formula(below.value):
        return None

    try:
        from_above = translator(above.value, origin=above.coordinate).translate_formula(target.coordinate)
        from_below = translator(below.value, origin=below.coordinate).translate_formula(target.coordinate)
    except Exception:
        # Structured references, external formulas, or unsupported formula
        # syntax must be reviewed rather than guessed.
        return None
    if from_above.startswith("=") and from_above == from_below:
        return from_above
    return None


def build_report(input_path: Path, output_path: Path, dry_run: bool) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "tool": "spreadsheet-formula-diagnostics",
        "created_at": utc_now(),
        "completed_at": None,
        "input_file": str(input_path.resolve()),
        "output_file": str(output_path.resolve()),
        "dry_run": dry_run,
        "summary": {
            "sheets_processed": 0,
            "formula_cells_found": 0,
            "formula_error_findings": 0,
            "broken_reference_findings": 0,
            "safe_formula_gap_candidates": 0,
            "safe_formula_gaps_repaired": 0,
        },
        "findings": [],
        "warnings": [],
    }


def load_workbooks(input_path: Path, keep_vba: bool) -> tuple[Any, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FormulaDiagnosticError(
            "Excel formula diagnostics require openpyxl. Install it in the active Python environment "
            "with: python3 -m pip install openpyxl."
        ) from exc
    try:
        formulas = load_workbook(input_path, data_only=False, keep_vba=keep_vba, keep_links=True)
        try:
            values = load_workbook(input_path, data_only=True, keep_vba=keep_vba, keep_links=True)
        except Exception:
            formulas.close()
            raise
        return formulas, values
    except Exception as exc:
        raise FormulaDiagnosticError(f"Could not open workbook for formula diagnostics: {exc}") from exc


def diagnose_workbook(
    input_path: Path,
    output_path: Path,
    *,
    report: dict[str, Any],
    selected_sheets: list[str] | None = None,
    include_hidden_sheets: bool = False,
    apply_safe_repairs: bool = False,
) -> dict[str, Any]:
    from openpyxl.formula.translate import Translator

    keep_vba = input_path.suffix.lower() == ".xlsm"
    formulas, values = load_workbooks(input_path, keep_vba)
    try:
        available_names = [sheet.title for sheet in formulas.worksheets]
        if selected_sheets:
            missing = [name for name in selected_sheets if name not in available_names]
            if missing:
                raise FormulaDiagnosticError(
                    f"Unknown worksheet(s): {', '.join(missing)}. Available: {', '.join(available_names)}"
                )
            targets = [formulas[name] for name in selected_sheets]
        else:
            targets = [
                sheet for sheet in formulas.worksheets
                if include_hidden_sheets or sheet.sheet_state == "visible"
            ]
        if not targets:
            raise FormulaDiagnosticError("No worksheets selected. Check --sheet or use --include-hidden-sheets.")

        for worksheet in targets:
            rectangle = worksheet.max_row * worksheet.max_column
            if rectangle > MAX_CELLS_PER_SHEET:
                raise FormulaDiagnosticError(
                    f"Worksheet {worksheet.title!r} spans {rectangle:,} cells, above the safety limit "
                    f"of {MAX_CELLS_PER_SHEET:,}. Split the workbook or select a smaller worksheet."
                )

        target_names = [sheet.title for sheet in targets]
        report["selected_sheets"] = target_names
        report["summary"]["sheets_processed"] = len(targets)
        if keep_vba:
            report["warnings"].append(
                "VBA preservation was enabled. Review advanced workbook features in Excel after saving the copy."
            )

        for sheet in formulas.worksheets:
            if sheet.title not in target_names:
                report["warnings"].append(f"Skipped worksheet {sheet.title!r} (not selected or hidden).")

        repairs: list[dict[str, Any]] = []
        for worksheet in targets:
            value_sheet = values[worksheet.title]
            for row in worksheet.iter_rows():
                for cell in row:
                    if is_formula(cell.value):
                        report["summary"]["formula_cells_found"] += 1
                        if "#REF!" in cell.value.upper():
                            report["summary"]["broken_reference_findings"] += 1
                            record(report, {
                                "kind": "broken_reference",
                                "severity": "error",
                                "sheet": worksheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "message": "Formula text contains a #REF! reference. The intended source cell cannot be inferred safely.",
                                "action": "manual_review",
                            })

                    cached = value_sheet[cell.coordinate]
                    if cell.data_type == "e" or cached.data_type == "e" or is_excel_error(cell.value) or is_excel_error(cached.value):
                        error_value = cell.value if cell.data_type == "e" or is_excel_error(cell.value) else cached.value
                        report["summary"]["formula_error_findings"] += 1
                        record(report, {
                            "kind": "excel_error_value",
                            "severity": "error",
                            "sheet": worksheet.title,
                            "cell": cell.coordinate,
                            "formula": cell.value if is_formula(cell.value) else None,
                            "value": error_value,
                            "cached_value": cached.value,
                            "message": "Cell contains an Excel error value. This tool reports the error but does not invent a replacement formula.",
                            "action": "manual_review",
                        })

            # Consider only one-cell vertical gaps with formulas immediately on
            # both sides. Every candidate must be independently implied by both.
            for row_index in range(2, worksheet.max_row):
                for column_index in range(1, worksheet.max_column + 1):
                    candidate = safe_gap_formula(worksheet, row_index, column_index, Translator)
                    if candidate is None:
                        continue
                    cell = worksheet.cell(row=row_index, column=column_index)
                    above = worksheet.cell(row=row_index - 1, column=column_index)
                    below = worksheet.cell(row=row_index + 1, column=column_index)
                    repair = {
                        "kind": "missing_formula_gap",
                        "severity": "warning",
                        "sheet": worksheet.title,
                        "cell": cell.coordinate,
                        "formula_above": above.value,
                        "formula_below": below.value,
                        "candidate_formula": candidate,
                        "safe_to_apply": True,
                        "action": "would_repair" if apply_safe_repairs else "suggested_only",
                        "message": "The formulas above and below translate to the same formula for this blank cell.",
                    }
                    report["summary"]["safe_formula_gap_candidates"] += 1
                    record(report, repair)
                    repairs.append(repair)

        if apply_safe_repairs and repairs:
            for repair in repairs:
                formulas[repair["sheet"]][repair["cell"]] = repair["candidate_formula"]
            report["summary"]["safe_formula_gaps_repaired"] = len(repairs)
            report["warnings"].append(
                "Formula repairs were written only to a new output copy. Review the copy in a compatible spreadsheet app to recalculate results."
            )
            calc = getattr(formulas, "calculation", None)
            if calc is not None:
                calc.fullCalcOnLoad = True
                calc.forceFullCalc = True
                calc.calcMode = "auto"

            expected_formula_count = formula_count(formulas, targets)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            fd, temp_name = tempfile.mkstemp(prefix=f".{output_path.stem}.", suffix=output_path.suffix, dir=str(output_path.parent))
            os.close(fd)
            try:
                formulas.save(temp_name)
                verification, verification_values = load_workbooks(Path(temp_name), keep_vba)
                try:
                    if verification.sheetnames != formulas.sheetnames:
                        raise FormulaDiagnosticError("Output validation failed: worksheet names changed.")
                    verification_targets = [verification[name] for name in target_names]
                    if formula_count(verification, verification_targets) != expected_formula_count:
                        raise FormulaDiagnosticError("Output validation failed: formula count changed in the selected worksheets.")
                    for repair in repairs:
                        actual = verification[repair["sheet"]][repair["cell"]].value
                        if actual != repair["candidate_formula"]:
                            raise FormulaDiagnosticError(
                                f"Output validation failed at {repair['sheet']}!{repair['cell']}."
                            )
                finally:
                    verification.close()
                    verification_values.close()
                os.replace(temp_name, output_path)
            finally:
                if os.path.exists(temp_name):
                    os.unlink(temp_name)
            report["output_written"] = True
        elif apply_safe_repairs:
            report["summary"]["safe_formula_gaps_repaired"] = 0
            report["warnings"].append(
                "No safely repairable single-cell formula gaps were found; no workbook copy was written."
            )
        else:
            report["warnings"].append(
                "Read-only diagnosis: no workbook copy was written. Use --apply-safe-repairs to create a copy containing only proven formula-gap repairs."
            )
        return report
    finally:
        formulas.close()
        values.close()


def write_report(path: Path, report: dict[str, Any], overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FormulaDiagnosticError(
            f"Report already exists: {path}. Choose a new --report path or pass --overwrite-output."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".json", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(report, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Diagnose spreadsheet formula issues and optionally repair only proven one-cell formula gaps."
    )
    parser.add_argument("input_file", type=Path, help="Input .xlsx or .xlsm workbook")
    parser.add_argument("--output", type=Path, help="Output copy path; defaults to <stem>_formulas_repaired.<ext>")
    parser.add_argument("--report", type=Path, help="JSON report path")
    parser.add_argument("--sheet", action="append", help="Worksheet to inspect; repeat to select multiple worksheets")
    parser.add_argument("--include-hidden-sheets", action="store_true", help="Include hidden sheets when no --sheet is given")
    parser.add_argument("--apply-safe-repairs", action="store_true", help="Write proven formula-gap repairs to a new output copy")
    parser.add_argument("--overwrite-output", action="store_true", help="Allow replacing existing output/report paths, never the input")
    return parser.parse_args(argv)


def run(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = args.input_file.expanduser().resolve()
    if not input_path.is_file():
        raise FormulaDiagnosticError(f"Input workbook does not exist or is not a regular file: {input_path}")
    suffix = input_path.suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise FormulaDiagnosticError("Supported formats are .xlsx and .xlsm. Convert legacy .xls files to .xlsx first.")

    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else input_path.with_name(input_path.stem + "_formulas_repaired" + input_path.suffix)
    )
    report_path = (
        args.report.expanduser().resolve()
        if args.report
        else input_path.with_name(input_path.stem + "_formula_diagnosis.json")
    )
    if output_path == input_path or report_path == input_path:
        raise FormulaDiagnosticError("Output and report paths must never overwrite the input file.")
    if output_path == report_path:
        raise FormulaDiagnosticError("The workbook output and JSON report must use different paths.")
    if output_path.suffix.lower() != suffix:
        raise FormulaDiagnosticError(f"Output extension must remain {suffix} to preserve the file type.")
    if args.apply_safe_repairs and output_path.exists() and not args.overwrite_output:
        raise FormulaDiagnosticError(f"Output already exists: {output_path}. Choose another path or use --overwrite-output.")
    if report_path.exists() and not args.overwrite_output:
        raise FormulaDiagnosticError(f"Report already exists: {report_path}. Choose another path or use --overwrite-output.")

    report = build_report(input_path, output_path, dry_run=not args.apply_safe_repairs)
    result = diagnose_workbook(
        input_path,
        output_path,
        report=report,
        selected_sheets=args.sheet,
        include_hidden_sheets=args.include_hidden_sheets,
        apply_safe_repairs=args.apply_safe_repairs,
    )
    result["completed_at"] = utc_now()
    write_report(report_path, result, overwrite=args.overwrite_output)
    print(json.dumps({
        "ok": True,
        "dry_run": not args.apply_safe_repairs,
        "input": str(input_path),
        "output": str(output_path) if args.apply_safe_repairs else None,
        "report": str(report_path),
        "summary": result["summary"],
        "warnings": result["warnings"],
    }, ensure_ascii=False, indent=2))
    return 0


def main() -> None:
    try:
        raise SystemExit(run())
    except FormulaDiagnosticError as exc:
        print(f"spreadsheet-formula-diagnostics: {exc}", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
