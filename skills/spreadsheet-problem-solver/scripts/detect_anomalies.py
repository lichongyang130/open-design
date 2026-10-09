#!/usr/bin/env python3
"""Detect missing values, exact duplicate rows, formula errors, and numeric outliers.

This is a read-only audit. Outliers are candidates for review, never assumed to
be business errors. The input file is never changed.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
import re
import sys
import tempfile
import unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
MAX_CELLS_PER_SHEET = 1_000_000
MAX_DATA_ROWS = 250_000
MAX_DETAILS = 5_000
ERROR_VALUES = {
    "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!",
    "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#UNKNOWN!", "#GETTING_DATA",
}
NUMERIC_TEXT = re.compile(r"^[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)$")


class AnomalyDetectionError(Exception):
    """Actionable input, safety, or validation failure."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def canonical_text(value: Any) -> str | None:
    if is_blank(value):
        return None
    return unicodedata.normalize("NFC", str(value)).strip()


def numeric_value(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        result = float(value)
        return result if math.isfinite(result) else None
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not NUMERIC_TEXT.fullmatch(text):
        return None
    # Text with a leading zero is often an identifier, not a quantity.
    # Keep it out of automatic numeric analysis unless it is a simple zero
    # or a fractional value such as 0.25.
    if re.match(r"^[+-]?0[0-9]+(?:\.[0-9]+)?$", text):
        return None
    try:
        result = float(Decimal(text))
    except (InvalidOperation, OverflowError, ValueError):
        return None
    return result if math.isfinite(result) else None


def add_detail(report: dict[str, Any], category: str, item: dict[str, Any]) -> None:
    details = report["findings"][category]
    if len(details) < MAX_DETAILS:
        details.append(item)
    else:
        report["details_truncated"] = True


def normalize_headers(raw: list[Any], file_label: str) -> list[str]:
    headers: list[str] = []
    for position, item in enumerate(raw, start=1):
        if is_blank(item):
            raise AnomalyDetectionError(
                f"{file_label} has an empty header in column {position}. Rename it before analysis."
            )
        name = unicodedata.normalize("NFC", str(item)).strip()
        if name in headers:
            raise AnomalyDetectionError(
                f"{file_label} has duplicate header {name!r}. Rename duplicate columns before analysis."
            )
        headers.append(name)
    return headers


def read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]], list[str], str | None]:
    raw = path.read_bytes()
    decoded: str | None = None
    encoding = "utf-8"
    for candidate in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise AnomalyDetectionError(f"Could not decode {path.name}. Export it as UTF-8 or GB18030 CSV.")
    try:
        delimiter = csv.Sniffer().sniff(decoded[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter))
    if not rows:
        raise AnomalyDetectionError(f"{path.name} is empty.")
    headers = normalize_headers(rows[0], path.name)
    if len(rows) - 1 > MAX_DATA_ROWS:
        raise AnomalyDetectionError(f"{path.name} exceeds the {MAX_DATA_ROWS:,}-row safety limit.")
    if len(rows) * len(headers) > MAX_CELLS_PER_SHEET:
        raise AnomalyDetectionError(f"{path.name} exceeds the {MAX_CELLS_PER_SHEET:,}-cell safety limit.")
    records: list[dict[str, Any]] = []
    for row_number, raw_row in enumerate(rows[1:], start=2):
        if len(raw_row) > len(headers) and any(not is_blank(v) for v in raw_row[len(headers):]):
            raise AnomalyDetectionError(
                f"{path.name} row {row_number} has non-empty values beyond the header columns."
            )
        values = raw_row[:len(headers)] + [None] * max(0, len(headers) - len(raw_row))
        records.append({"row_number": row_number, "values": dict(zip(headers, values))})
    return headers, records, [f"CSV delimiter: {repr(delimiter)}; source encoding decoded as {encoding}."], None


def read_excel(path: Path, sheet_name: str | None) -> tuple[list[str], list[dict[str, Any]], list[str], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AnomalyDetectionError(
            "Excel anomaly detection requires openpyxl. Install it with: python3 -m pip install openpyxl."
        ) from exc
    try:
        value_book = load_workbook(path, data_only=True, read_only=True, keep_vba=path.suffix.lower() == ".xlsm")
        formula_book = load_workbook(path, data_only=False, read_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    except Exception as exc:
        raise AnomalyDetectionError(f"Could not open workbook {path.name}: {exc}") from exc
    try:
        if sheet_name:
            if sheet_name not in value_book.sheetnames:
                raise AnomalyDetectionError(
                    f"Worksheet {sheet_name!r} was not found in {path.name}. Available: {', '.join(value_book.sheetnames)}"
                )
            sheet = value_book[sheet_name]
            formula_sheet = formula_book[sheet_name]
        else:
            visible = [ws.title for ws in value_book.worksheets if ws.sheet_state == "visible"]
            if not visible:
                raise AnomalyDetectionError(f"{path.name} has no visible worksheet.")
            sheet = value_book[visible[0]]
            formula_sheet = formula_book[visible[0]]
            sheet_name = visible[0]
        rectangle = sheet.max_row * sheet.max_column
        if rectangle > MAX_CELLS_PER_SHEET:
            raise AnomalyDetectionError(
                f"{path.name}:{sheet.title} exceeds the {MAX_CELLS_PER_SHEET:,}-cell safety limit."
            )
        if max(0, sheet.max_row - 1) > MAX_DATA_ROWS:
            raise AnomalyDetectionError(f"{path.name}:{sheet.title} exceeds the {MAX_DATA_ROWS:,}-row safety limit.")
        raw_rows = list(sheet.iter_rows(values_only=True))
        formula_rows = list(formula_sheet.iter_rows(values_only=False))
        if not raw_rows:
            raise AnomalyDetectionError(f"{path.name}:{sheet.title} is empty.")
        headers = normalize_headers(list(raw_rows[0]), f"{path.name}:{sheet.title}")
        records = []
        warnings: list[str] = []
        for row_number, (raw_row, formula_row) in enumerate(zip(raw_rows[1:], formula_rows[1:]), start=2):
            if len(raw_row) > len(headers) and any(not is_blank(v) for v in raw_row[len(headers):]):
                raise AnomalyDetectionError(
                    f"{path.name}:{sheet.title} row {row_number} has values beyond the header columns."
                )
            values = list(raw_row[:len(headers)]) + [None] * max(0, len(headers) - len(raw_row))
            formula_cells = list(formula_row[:len(headers)])
            formula_values = [cell.value for cell in formula_cells] + [None] * max(0, len(headers) - len(formula_cells))
            records.append({
                "row_number": row_number,
                "values": dict(zip(headers, values)),
                "formula_values": dict(zip(headers, formula_values)),
            })
            for position, formula_cell in enumerate(formula_row[:len(headers)]):
                if formula_cell.data_type == "f" and position < len(raw_row) and raw_row[position] is None:
                    warnings.append(
                        f"{path.name}:{sheet.title}!{formula_cell.coordinate} has a formula with no cached value; recalculate and save the workbook before relying on numeric analysis."
                    )
        return headers, records, list(dict.fromkeys(warnings)), sheet.title
    finally:
        value_book.close()
        formula_book.close()


def percentile(sorted_values: list[float], fraction: float) -> float:
    if len(sorted_values) == 1:
        return sorted_values[0]
    position = (len(sorted_values) - 1) * fraction
    left = math.floor(position)
    right = math.ceil(position)
    if left == right:
        return sorted_values[left]
    weight = position - left
    return sorted_values[left] * (1 - weight) + sorted_values[right] * weight


def error_text(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() in ERROR_VALUES


def detect(
    input_path: Path,
    *,
    sheet_name: str | None = None,
    required_columns: list[str] | None = None,
    numeric_columns: list[str] | None = None,
    auto_numeric: bool = True,
    iqr_multiplier: float = 1.5,
) -> dict[str, Any]:
    suffix = input_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise AnomalyDetectionError("Supported formats are CSV, XLSX, and XLSM. Convert old XLS files to XLSX first.")
    if suffix == ".csv":
        headers, records, warnings, selected_sheet = read_csv(input_path)
    else:
        headers, records, warnings, selected_sheet = read_excel(input_path, sheet_name)
    required_columns = required_columns or []
    numeric_columns = numeric_columns or []
    missing_required = [column for column in required_columns if column not in headers]
    missing_numeric = [column for column in numeric_columns if column not in headers]
    if missing_required:
        raise AnomalyDetectionError(f"Required column(s) not found: {', '.join(missing_required)}.")
    if missing_numeric:
        raise AnomalyDetectionError(f"Numeric column(s) not found: {', '.join(missing_numeric)}.")
    if not math.isfinite(iqr_multiplier) or iqr_multiplier <= 0:
        raise AnomalyDetectionError("--iqr-multiplier must be a finite number greater than zero.")

    report: dict[str, Any] = {
        "schema_version": 1,
        "tool": "spreadsheet-anomaly-detection",
        "created_at": utc_now(),
        "completed_at": None,
        "input_file": str(input_path.resolve()),
        "selected_sheet": selected_sheet,
        "parameters": {
            "required_columns": required_columns,
            "numeric_columns": numeric_columns,
            "auto_numeric": auto_numeric,
            "iqr_multiplier": iqr_multiplier,
        },
        "summary": {
            "rows_scanned": len(records),
            "missing_values": 0,
            "error_values": 0,
            "broken_reference_formulas": 0,
            "exact_duplicate_rows": 0,
            "numeric_columns_analyzed": [],
            "numeric_outlier_candidates": 0,
        },
        "findings": {
            "missing_values": [],
            "error_values": [],
            "broken_references": [],
            "exact_duplicate_rows": [],
            "numeric_outliers": [],
        },
        "warnings": warnings,
    }

    columns_to_check = required_columns or headers
    for record in records:
        values = record["values"]
        formula_values = record.get("formula_values", {})
        for column in columns_to_check:
            has_formula = isinstance(formula_values.get(column), str) and formula_values[column].startswith("=")
            # Excel formula cells may have an empty cached result before the
            # workbook has been recalculated; do not mislabel those as blanks.
            if is_blank(values.get(column)) and not has_formula:
                report["summary"]["missing_values"] += 1
                add_detail(report, "missing_values", {
                    "row": record["row_number"], "column": column,
                    "cell_value": json_value(values.get(column)),
                    "issue": "blank_required_value" if required_columns else "blank_cell",
                })
        for column in headers:
            value = values.get(column)
            formula_value = formula_values.get(column, value)
            if error_text(value) or (isinstance(value, str) and value.strip().upper() in ERROR_VALUES):
                report["summary"]["error_values"] += 1
                add_detail(report, "error_values", {
                    "row": record["row_number"], "column": column, "value": json_value(value),
                    "issue": "spreadsheet_error_value",
                })
            if isinstance(formula_value, str) and formula_value.startswith("=") and "#REF!" in formula_value.upper():
                report["summary"]["broken_reference_formulas"] += 1
                add_detail(report, "broken_references", {
                    "row": record["row_number"], "column": column,
                    "formula": formula_value, "issue": "formula_contains_ref_error",
                })

    seen_rows: dict[tuple[str | None, ...], int] = {}
    for record in records:
        values = record["values"]
        signature = tuple(canonical_text(values.get(column)) for column in headers)
        if all(part is None for part in signature):
            continue
        previous_row = seen_rows.get(signature)
        if previous_row is None:
            seen_rows[signature] = record["row_number"]
            continue
        report["summary"]["exact_duplicate_rows"] += 1
        add_detail(report, "exact_duplicate_rows", {
            "row": record["row_number"], "kept_row": previous_row,
            "issue": "exact_duplicate_row", "action": "review_only",
        })

    selected_numeric = list(numeric_columns)
    if not selected_numeric and auto_numeric:
        for column in headers:
            samples = [numeric_value(record["values"].get(column)) for record in records]
            nonblank = [record["values"].get(column) for record in records if not is_blank(record["values"].get(column))]
            numeric_samples = [value for value in samples if value is not None]
            if len(numeric_samples) >= 4 and len(numeric_samples) / max(1, len(nonblank)) >= 0.8:
                selected_numeric.append(column)
    for column in selected_numeric:
        indexed: list[tuple[int, float, Any]] = []
        for record in records:
            original = record["values"].get(column)
            parsed = numeric_value(original)
            if parsed is not None:
                indexed.append((record["row_number"], parsed, original))
        if len(indexed) < 4:
            report["warnings"].append(
                f"Column {column!r} has fewer than four parseable numeric values; IQR outlier detection was skipped."
            )
            continue
        sorted_values = sorted(value for _, value, _ in indexed)
        q1 = percentile(sorted_values, 0.25)
        q3 = percentile(sorted_values, 0.75)
        spread = q3 - q1
        lower = q1 - iqr_multiplier * spread
        upper = q3 + iqr_multiplier * spread
        report["summary"]["numeric_columns_analyzed"].append(column)
        if spread == 0:
            report["warnings"].append(
                f"Column {column!r} has an interquartile range of zero; statistical outlier candidates were not emitted."
            )
            continue
        for row_number, value, original in indexed:
            if value < lower or value > upper:
                report["summary"]["numeric_outlier_candidates"] += 1
                add_detail(report, "numeric_outliers", {
                    "row": row_number, "column": column,
                    "original_value": json_value(original), "numeric_value": value,
                    "q1": q1, "q3": q3, "iqr": spread,
                    "lower_bound": lower, "upper_bound": upper,
                    "rule": f"IQR {iqr_multiplier:g}x fence",
                    "issue": "statistical_outlier_candidate",
                    "action": "review_only",
                })
    report["summary"]["details_truncated"] = bool(report.get("details_truncated"))
    return report


def write_json(path: Path, report: dict[str, Any], overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise AnomalyDetectionError(f"Report already exists: {path}. Choose another --report path or use --overwrite-output.")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".json", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(report, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read-only anomaly audit for CSV/XLSX/XLSM with JSON report output."
    )
    parser.add_argument("input_file", type=Path)
    parser.add_argument("--sheet", help="Worksheet name in Excel; defaults to the first visible sheet")
    parser.add_argument("--required-column", action="append", help="Column where blanks should be flagged as missing-required values")
    parser.add_argument("--numeric-column", action="append", help="Numeric column for IQR outlier detection; repeat as needed")
    parser.add_argument("--no-auto-numeric", action="store_true", help="Do not automatically scan numeric-looking columns")
    parser.add_argument("--iqr-multiplier", type=float, default=1.5, help="IQR fence multiplier; default 1.5")
    parser.add_argument("--report", type=Path, help="JSON report path")
    parser.add_argument("--overwrite-output", action="store_true", help="Allow replacing the report, never the input")
    return parser.parse_args(argv)


def run(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = args.input_file.expanduser().resolve()
    if not input_path.is_file():
        raise AnomalyDetectionError(f"Input file does not exist or is not a regular file: {input_path}")
    report_path = (
        args.report.expanduser().resolve()
        if args.report
        else input_path.with_name(input_path.stem + "_anomaly_report.json")
    )
    if report_path == input_path:
        raise AnomalyDetectionError("The JSON report must never overwrite the input file.")
    if report_path.exists() and not args.overwrite_output:
        raise AnomalyDetectionError(f"Report already exists: {report_path}. Choose another path or use --overwrite-output.")
    report = detect(
        input_path,
        sheet_name=args.sheet,
        required_columns=args.required_column,
        numeric_columns=args.numeric_column,
        auto_numeric=not args.no_auto_numeric,
        iqr_multiplier=args.iqr_multiplier,
    )
    report["completed_at"] = utc_now()
    write_json(report_path, report, overwrite=args.overwrite_output)
    print(json.dumps({
        "ok": True,
        "input": str(input_path),
        "report": str(report_path),
        "summary": report["summary"],
        "warnings": report["warnings"],
    }, ensure_ascii=False, indent=2))
    return 0


def main() -> None:
    try:
        raise SystemExit(run())
    except AnomalyDetectionError as exc:
        print(f"detect_spreadsheet_anomalies: {exc}", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
