from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import clean_spreadsheet as cleaner  # noqa: E402


class CsvCleaningTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-spreadsheet-cleaner-")
        self.root = Path(self.temp.name)
        self.input_path = self.root / "sales.csv"
        self.output_path = self.root / "sales_cleaned.csv"

    def tearDown(self) -> None:
        self.temp.cleanup()

    def read_output(self) -> list[list[str]]:
        with self.output_path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.reader(handle))

    def test_default_trims_text_but_only_reports_blank_and_duplicate_rows(self) -> None:
        self.input_path.write_text(
            "Order,Amount,Code\n A-1 , 12 ,00123\n\nA-1,12,00123\n",
            encoding="utf-8",
        )
        report = cleaner._build_report(self.input_path, self.output_path, "csv", False)

        cleaner.clean_csv(self.input_path, self.output_path, report=report)

        self.assertEqual(
            self.read_output(),
            [
                ["Order", "Amount", "Code"],
                ["A-1", "12", "00123"],
                [],
                ["A-1", "12", "00123"],
            ],
        )
        self.assertEqual(report["summary"]["text_cells_trimmed"], 2)
        self.assertEqual(report["summary"]["blank_rows_detected"], 1)
        self.assertEqual(report["summary"]["blank_rows_removed"], 0)
        self.assertEqual(report["summary"]["duplicate_rows_detected"], 1)
        self.assertEqual(report["summary"]["duplicate_rows_removed"], 0)
        self.assertTrue(self.input_path.exists(), "source file must remain intact")

    def test_explicit_options_drop_blank_rows_and_exact_duplicates(self) -> None:
        self.input_path.write_text(
            "Name,Amount\nAlice,100\n\nBob,50\nAlice,100\n",
            encoding="utf-8",
        )
        report = cleaner._build_report(self.input_path, self.output_path, "csv", False)

        cleaner.clean_csv(
            self.input_path,
            self.output_path,
            report=report,
            drop_empty_rows=True,
            dedupe=True,
        )

        self.assertEqual(
            self.read_output(),
            [["Name", "Amount"], ["Alice", "100"], ["Bob", "50"]],
        )
        self.assertEqual(report["summary"]["blank_rows_removed"], 1)
        self.assertEqual(report["summary"]["duplicate_rows_detected"], 1)
        self.assertEqual(report["summary"]["duplicate_rows_removed"], 1)

    def test_dry_run_does_not_write_cleaned_file(self) -> None:
        self.input_path.write_text("Name\n Alice \n", encoding="utf-8")
        report = cleaner._build_report(self.input_path, self.output_path, "csv", True)

        cleaner.clean_csv(self.input_path, self.output_path, report=report, dry_run=True)

        self.assertFalse(self.output_path.exists())
        self.assertEqual(report["summary"]["text_cells_planned_to_trim"], 1)
        self.assertEqual(report["summary"]["text_cells_trimmed"], 0)
        self.assertTrue(self.input_path.exists())

    def test_dry_run_reports_planned_deletions_without_claiming_applied_changes(self) -> None:
        self.input_path.write_text("Name,Amount\nAlice,100\n\nAlice,100\n", encoding="utf-8")
        report = cleaner._build_report(self.input_path, self.output_path, "csv", True)

        cleaner.clean_csv(
            self.input_path,
            self.output_path,
            report=report,
            drop_empty_rows=True,
            dedupe=True,
            dry_run=True,
        )

        self.assertFalse(self.output_path.exists())
        self.assertEqual(report["summary"]["blank_rows_removed"], 0)
        self.assertEqual(report["summary"]["blank_rows_planned_for_removal"], 1)
        self.assertEqual(report["summary"]["duplicate_rows_removed"], 0)
        self.assertEqual(report["summary"]["duplicate_rows_planned_for_removal"], 1)
        self.assertEqual(report["summary"]["rows_after_planned_changes"], 2)
        self.assertNotIn("rows_written", report["summary"])
        self.assertTrue(self.input_path.exists())

    def test_multiline_csv_fields_survive_round_trip(self) -> None:
        self.input_path.write_text('Name,Notes\nAlice,"first line\nsecond line"\n', encoding="utf-8")
        report = cleaner._build_report(self.input_path, self.output_path, "csv", False)

        cleaner.clean_csv(self.input_path, self.output_path, report=report)

        self.assertEqual(self.read_output(), [["Name", "Notes"], ["Alice", "first line\nsecond line"]])

    def test_formula_like_text_is_not_exposed_by_trimming(self) -> None:
        self.input_path.write_text('Name,Value\n" =1+1 ",ordinary\n', encoding="utf-8")
        report = cleaner._build_report(self.input_path, self.output_path, "csv", False)

        cleaner.clean_csv(self.input_path, self.output_path, report=report)

        self.assertEqual(self.read_output(), [["Name", "Value"], [" =1+1 ", "ordinary"]])
        self.assertEqual(report["summary"]["text_cells_trimmed"], 0)
        self.assertTrue(
            any(finding.get("kind") == "formula_like_text_preserved" for finding in report["findings"])
        )

    def test_source_file_cannot_be_selected_as_output(self) -> None:
        self.input_path.write_text("Name\nAlice\n", encoding="utf-8")
        report_path = self.root / "report.json"

        with self.assertRaises(cleaner.CleanerError):
            cleaner.run([str(self.input_path), "--output", str(self.input_path), "--report", str(report_path)])

        self.assertTrue(self.input_path.exists())
        self.assertFalse(report_path.exists())


try:
    import openpyxl  # noqa: F401
except ImportError:
    openpyxl = None


@unittest.skipIf(openpyxl is None, "openpyxl is optional; Excel tests require it")
class WorkbookCleaningTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="od-spreadsheet-cleaner-xlsx-")
        self.root = Path(self.temp.name)
        self.input_path = self.root / "sales.xlsx"
        self.output_path = self.root / "sales_cleaned.xlsx"

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_trims_text_preserves_formula_and_sheet_names(self) -> None:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Sales"
        worksheet.append(["Name", "Total"])
        worksheet.append([" Alice ", "=1+2"])
        worksheet.append([" ", None])  # A whitespace-only row must be detected after trimming.
        workbook.save(self.input_path)

        report = cleaner._build_report(self.input_path, self.output_path, "xlsx", False)
        cleaner.clean_workbook(self.input_path, self.output_path, report=report)

        output = openpyxl.load_workbook(self.output_path, data_only=False)
        try:
            self.assertEqual(output.sheetnames, ["Sales"])
            self.assertEqual(output["Sales"]["A2"].value, "Alice")
            self.assertEqual(output["Sales"]["B2"].value, "=1+2")
        finally:
            output.close()
        self.assertEqual(report["summary"]["text_cells_trimmed"], 2)
        self.assertEqual(report["summary"]["formula_cells_preserved"], 1)
        self.assertEqual(report["summary"]["blank_rows_detected"], 1)

    def test_opt_in_dedupe_clears_values_without_shifting_rows_or_formulas(self) -> None:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["Order", "Amount", "Check"])
        worksheet.append(["A-1", 100, None])
        worksheet.append(["A-1", 100, None])
        worksheet.append(["Total", None, "=SUM(B2:B3)"])
        workbook.save(self.input_path)

        report = cleaner._build_report(self.input_path, self.output_path, "xlsx", False)
        cleaner.clean_workbook(self.input_path, self.output_path, report=report, dedupe=True)

        output = openpyxl.load_workbook(self.output_path, data_only=False)
        try:
            sheet = output.active
            self.assertEqual(sheet["A2"].value, "A-1")
            self.assertIsNone(sheet["A3"].value)
            self.assertIsNone(sheet["B3"].value)
            self.assertEqual(sheet["C4"].value, "=SUM(B2:B3)")
        finally:
            output.close()
        self.assertEqual(report["summary"]["duplicate_rows_detected"], 1)
        self.assertEqual(report["summary"]["duplicate_rows_cleared"], 1)

    def test_formula_bearing_duplicate_row_is_not_cleared(self) -> None:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["Description", "Formula"])
        worksheet.append(["Calculation", "=1+2"])
        worksheet.append(["Calculation", "=1+2"])
        workbook.save(self.input_path)

        report = cleaner._build_report(self.input_path, self.output_path, "xlsx", False)
        cleaner.clean_workbook(self.input_path, self.output_path, report=report, dedupe=True)

        output = openpyxl.load_workbook(self.output_path, data_only=False)
        try:
            self.assertEqual(output.active["A3"].value, "Calculation")
            self.assertEqual(output.active["B3"].value, "=1+2")
        finally:
            output.close()
        self.assertEqual(report["summary"]["duplicate_rows_cleared"], 0)
        self.assertTrue(
            any(
                finding.get("action") == "not_cleared_formula_cells_present"
                for finding in report["findings"]
            )
        )


if __name__ == "__main__":
    unittest.main()
